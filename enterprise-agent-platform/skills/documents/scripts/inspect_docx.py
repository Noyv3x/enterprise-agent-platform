#!/usr/bin/env python3
"""Structural DOCX check. Usage: inspect_docx.py FILE.docx
Prints headings, paragraph/table counts, section geometry, runs lacking eastAsia fonts that contain
CJK text, tables wider than the text area. Does NOT prove pagination or appearance: use `artifact view`."""
import re, sys
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu

d = Document(sys.argv[1])
s = d.sections[0]
text_w = s.page_width - s.left_margin - s.right_margin
print(f"paragraphs={len(d.paragraphs)} tables={len(d.tables)} sections={len(d.sections)} "
      f"page={Emu(s.page_width).cm:.1f}x{Emu(s.page_height).cm:.1f}cm text_width={Emu(text_w).cm:.1f}cm "
      f"inline_images={len(d.inline_shapes)}")
for p in d.paragraphs:
    if p.style.name.startswith(("Heading", "Title")):
        print(f"  [{p.style.name}] {p.text[:60]}")
cjk = re.compile(r"[\u3400-\u9fff]")
def east_font(style):
    while style is not None:
        rpr = style.element.rPr
        f = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if f is not None and f.get(qn("w:eastAsia")):
            return True
        style = style.base_style
    return False
missing = 0
for p in d.paragraphs:
    if cjk.search(p.text):
        for r in p.runs:
            rpr = r._r.rPr; f = rpr.find(qn("w:rFonts")) if rpr is not None else None
            if cjk.search(r.text) and not (f is not None and f.get(qn("w:eastAsia"))) and not east_font(p.style):
                missing += 1
print(f"cjk runs without an eastAsia font (run or style): {missing}")
for i, t in enumerate(d.tables, 1):
    w = sum(c.width or 0 for c in t.rows[0].cells)
    flag = "  <-- WIDER THAN TEXT AREA" if w > text_w * 1.01 else ""
    print(f"  table {i}: {len(t.rows)}x{len(t.columns)} width={Emu(w).cm:.1f}cm{flag}")
