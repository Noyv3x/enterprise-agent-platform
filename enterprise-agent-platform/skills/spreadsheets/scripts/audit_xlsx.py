#!/usr/bin/env python3
"""Audit an XLSX. Usage: audit_xlsx.py FILE.xlsx
Per sheet: dimensions, formulas, formulas WITHOUT a cached value, Excel error values (#REF! ...),
merged cells, charts/images, hard-coded numbers sitting inside formulas, and a sample.
Exit 1 when errors or missing cached values are found.
It reads values cached by whoever saved the file; it does not recalculate anything."""
import re, sys
import openpyxl

path = sys.argv[1]
f = openpyxl.load_workbook(path)                    # formulas as text
v = openpyxl.load_workbook(path, data_only=True)    # cached values
ERR = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
problems = 0
print("defined names:", {d.name: d.attr_text for d in f.defined_names.definedName})
for ws in f.worksheets:
    vs = v[ws.title]; nform = nocache = errs = 0; hard = []
    for row in ws.iter_rows():
        for c in row:
            val = c.value
            cached = vs[c.coordinate].value
            if isinstance(val, str) and val.startswith("="):
                nform += 1
                if cached is None:
                    nocache += 1
                if re.search(r"[*/+\-^]\s*\d+\.\d+|[*/+\-^]\s*\d{3,}", val):
                    hard.append(c.coordinate)  # constants buried in formulas belong on the assumptions sheet
            if isinstance(cached, str) and cached.strip() in ERR:
                errs += 1; print(f"  ERROR {ws.title}!{c.coordinate} = {cached}")
    problems += errs + nocache
    print(f"[{ws.title}] dims={ws.dimensions} formulas={nform} no_cached_value={nocache} errors={errs} "
          f"merged={len(ws.merged_cells.ranges)} charts={len(ws._charts)} images={len(ws._images)} state={ws.sheet_state}"
          + (f" suspicious_hardcodes={hard[:8]}" if hard else ""))
    for r in list(vs.iter_rows(min_row=1, max_row=4, values_only=True)):
        print("   ", [x for x in r][:8])
if problems:
    print("PROBLEMS: formulas without cached values or error values (preview/pandas will show blanks or errors)")
    sys.exit(1)
