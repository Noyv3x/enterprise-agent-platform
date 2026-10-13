"""python-pptx (1.0.2) helpers: 16:9 deck, real title placeholders, CJK fonts, native charts/tables.

    import sys; sys.path.insert(0, "/opt/agent-platform/skills/builtin/presentations/scripts")
    from pptx_kit import *

Every slide is built from the 'Title Only' layout so titles are real title placeholders
(outline view, accessibility, slide sorter); everything else is positioned on a 12-column grid.
"""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_W = SLIDE_W - 2 * MARGIN
LATIN, EAST = "Calibri", "Microsoft YaHei"
INK, MUTED, ACCENT, TINT, WHITE = (RGBColor.from_string(c) for c in ("1F2933", "52606D", "1D4ED8", "EEF2F7", "FFFFFF"))
PALETTE = [RGBColor.from_string(c) for c in ("1D4ED8", "F59E0B", "059669", "7C3AED", "DC2626", "64748B")]
SIZE = {"title": 30, "body": 20, "small": 14, "caption": 12}   # pt; never go below 12


def set_font(run, size, bold=False, color=INK, latin=LATIN, east=EAST):
    f = run.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = latin
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", east)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def add_slide(prs, title, notes=None):
    """Title Only layout; the title placeholder is moved/styled to the grid."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = MARGIN, Inches(0.45), CONTENT_W, Inches(1.0)
    tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = title
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    set_font(p.runs[0], SIZE["title"], bold=True)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_text(slide, left, top, width, height, paragraphs, size=SIZE["body"], color=INK, bullets=False, bold_first=False, anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of str. Bullets use a real bullet character (a:buChar) with hanging indent."""
    box = slide.shapes.add_textbox(left, top, width, height); tf = box.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, text in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = text
        set_font(r, size, bold=bold_first and i == 0, color=color)
        p.space_after = Pt(size * 0.5); p.line_spacing = 1.15
        if bullets:
            ppr = p._p.get_or_add_pPr(); ppr.set("marL", str(int(Pt(size * 1.2)))); ppr.set("indent", str(-int(Pt(size * 1.2))))
            bu = etree.SubElement(ppr, qn("a:buChar")); bu.set("char", "•")
    return box


def add_card(slide, left, top, width, height, heading, body, accent=ACCENT):
    """Tinted card: accent bar + heading + body (a KPI or a one-line argument)."""
    card = slide.shapes.add_shape(1, left, top, width, height)   # 1 = rectangle
    card.fill.solid(); card.fill.fore_color.rgb = TINT; card.line.fill.background(); card.shadow.inherit = False
    bar = slide.shapes.add_shape(1, left, top, Inches(0.08), height)
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background(); bar.shadow.inherit = False
    add_text(slide, left + Inches(0.25), top + Inches(0.15), width - Inches(0.4), Inches(0.7), [heading], size=SIZE["title"] - 4, color=accent, bold_first=True)
    add_text(slide, left + Inches(0.25), top + Inches(0.95), width - Inches(0.4), height - Inches(1.1), [body], size=SIZE["small"], color=MUTED)


def add_chart(slide, left, top, width, height, categories, series, kind=XL_CHART_TYPE.COLUMN_CLUSTERED, number_format='#,##0', legend=True):
    """Native, editable chart. series = {"name": [values...]}."""
    data = CategoryChartData(); data.categories = categories
    for name, values in series.items():
        data.add_series(name, values)
    gf = slide.shapes.add_chart(kind, left, top, width, height, data); ch = gf.chart
    ch.font.size = Pt(SIZE["small"]); ch.font.name = LATIN
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    plot = ch.plots[0]; plot.has_data_labels = True
    plot.data_labels.number_format = number_format; plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(SIZE["small"])
    if hasattr(plot, "gap_width"):
        plot.gap_width = 80
    if hasattr(plot, "overlap"):
        plot.overlap = -10
    for i, s in enumerate(plot.series):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = PALETTE[i % len(PALETTE)]
    return ch


def add_table(slide, left, top, width, rows, col_widths_in, row_h=Inches(0.5), size=SIZE["small"], number_cols=()):
    gt = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, row_h * len(rows)); t = gt.table
    for c, w in enumerate(col_widths_in):
        t.columns[c].width = Inches(w)
    for r, row in enumerate(rows):
        t.rows[r].height = row_h
        for c, value in enumerate(row):
            cell = t.cell(r, c); cell.text = str(value); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT if r == 0 else (TINT if r % 2 == 0 else WHITE)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if c in number_cols else PP_ALIGN.LEFT
            set_font(p.runs[0], size, bold=r == 0, color=WHITE if r == 0 else INK)
    return t


def add_footer(slide, text, page):
    add_text(slide, MARGIN, SLIDE_H - Inches(0.55), Inches(9), Inches(0.35), [text], size=SIZE["caption"], color=MUTED)
    add_text(slide, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.55), Inches(1), Inches(0.35), [str(page)], size=SIZE["caption"], color=MUTED)


def add_picture_fit(slide, path, left, top, max_w, max_h):
    """Insert an image scaled to fit the box, aspect ratio preserved, centered."""
    from PIL import Image
    w, h = Image.open(path).size
    scale = min(max_w / w, max_h / h)
    pw, ph = int(w * scale), int(h * scale)
    return slide.shapes.add_picture(path, left + (max_w - pw) // 2, top + (max_h - ph) // 2, pw, ph)
